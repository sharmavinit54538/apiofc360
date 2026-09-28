"""Document Generator Service for AI Hub Gateway.

Renders document templates with dynamic variable substitution, produces PDF, DOCX,
HTML, and TXT artifacts, and persists generated documents with tenant scoping.
"""

from __future__ import annotations

from datetime import datetime, timezone
import logging
import os
import re
from typing import Any, Optional
import uuid

from sqlalchemy import desc, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import AppException, NotFoundException
from app.models.ai_hub import GeneratedDocument
from app.models.document.template import DocumentTemplate
from app.schemas.ai_hub.document_generator import (
    DocumentGeneratorOverview,
    DocumentTemplateItem,
    DocumentTemplatesList,
    GenerateDocumentRequest,
    GenerateDocumentResponse,
    PreviewDocumentRequest,
    PreviewDocumentResponse,
)

logger = logging.getLogger(__name__)

DOCS_OUTPUT_DIR = os.path.join("uploads", "generated_documents")


class DocumentGeneratorService:
    """Service handling template discovery, preview rendering, and document artifact generation."""

    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        os.makedirs(DOCS_OUTPUT_DIR, exist_ok=True)

    @staticmethod
    def extract_placeholders(template_body: str) -> list[str]:
        """Extract {{variable}} and {variable} placeholders from template text."""
        placeholders = set()
        # Double curly braces {{foo}}
        for m in re.finditer(r"\{\{([a-zA-Z0-9_]+)\}\}", template_body):
            placeholders.add(m.group(1))
        # Single curly braces {foo}
        for m in re.finditer(r"\{([a-zA-Z0-9_]+)\}", template_body):
            placeholders.add(m.group(1))
        return sorted(list(placeholders))

    @staticmethod
    def render_template_body(template_body: str, variables: dict[str, Any]) -> str:
        """Substitute placeholders with provided variable values."""
        rendered = template_body
        for k, v in variables.items():
            str_val = str(v) if v is not None else ""
            rendered = rendered.replace(f"{{{{{k}}}}}", str_val)
            rendered = rendered.replace(f"{{{k}}}", str_val)
        return rendered

    async def get_overview(
        self, company_id: Optional[uuid.UUID] = None
    ) -> DocumentGeneratorOverview:
        """Fetch summary stats for the document generator engine."""
        # Total templates
        stmt_templates = select(func.count(DocumentTemplate.id))
        res_t = await self.session.execute(stmt_templates)
        total_templates = res_t.scalar() or 0

        # Total generated documents for company
        stmt_docs = select(func.count(GeneratedDocument.id))
        if company_id is not None:
            stmt_docs = stmt_docs.where(GeneratedDocument.company_id == company_id)
        res_d = await self.session.execute(stmt_docs)
        total_generated = res_d.scalar() or 0

        # Recent generated
        stmt_recent = (
            select(GeneratedDocument)
            .order_by(desc(GeneratedDocument.created_at))
            .limit(5)
        )
        if company_id is not None:
            stmt_recent = stmt_recent.where(GeneratedDocument.company_id == company_id)
        res_recent = await self.session.execute(stmt_recent)
        recent_docs = res_recent.scalars().all()

        recent_items = [
            {
                "documentId": str(doc.id),
                "title": doc.title,
                "format": doc.format,
                "createdAt": doc.created_at.isoformat(),
            }
            for doc in recent_docs
        ]

        return DocumentGeneratorOverview(
            totalTemplates=total_templates,
            totalGeneratedDocuments=total_generated,
            supportedFormats=["pdf", "docx", "html", "txt"],
            recentGenerated=recent_items,
        )

    async def list_templates(self) -> DocumentTemplatesList:
        """List all document templates available for document generation."""
        stmt = select(DocumentTemplate).order_by(DocumentTemplate.name.asc())
        res = await self.session.execute(stmt)
        templates = res.scalars().all()

        items = [
            DocumentTemplateItem(
                templateId=str(t.id),
                name=t.name,
                description=t.description,
                placeholders=self.extract_placeholders(t.template_body),
                createdAt=t.created_at.isoformat(),
            )
            for t in templates
        ]

        return DocumentTemplatesList(total=len(items), templates=items)

    async def preview(self, payload: PreviewDocumentRequest) -> PreviewDocumentResponse:
        """Render a preview of a template with substituted variables without persistence."""
        try:
            template_uuid = uuid.UUID(payload.templateId)
        except (ValueError, TypeError):
            raise AppException(message=f"Invalid templateId format: '{payload.templateId}'.")

        stmt = select(DocumentTemplate).where(DocumentTemplate.id == template_uuid)
        res = await self.session.execute(stmt)
        template = res.scalars().first()
        if not template:
            raise NotFoundException(message=f"Document template '{payload.templateId}' not found.")

        rendered = self.render_template_body(template.template_body, payload.variables)
        remaining_placeholders = self.extract_placeholders(rendered)

        return PreviewDocumentResponse(
            templateId=str(template.id),
            templateName=template.name,
            renderedContent=rendered,
            missingPlaceholders=remaining_placeholders,
            variablesUsed=payload.variables,
        )

    async def generate(
        self,
        company_id: Optional[uuid.UUID],
        user_id: Optional[uuid.UUID],
        payload: GenerateDocumentRequest,
    ) -> GenerateDocumentResponse:
        """Render template, generate physical document artifact, and persist record."""
        try:
            template_uuid = uuid.UUID(payload.templateId)
        except (ValueError, TypeError):
            raise AppException(message=f"Invalid templateId format: '{payload.templateId}'.")

        stmt = select(DocumentTemplate).where(DocumentTemplate.id == template_uuid)
        res = await self.session.execute(stmt)
        template = res.scalars().first()
        if not template:
            raise NotFoundException(message=f"Document template '{payload.templateId}' not found.")

        rendered_text = self.render_template_body(template.template_body, payload.variables)
        doc_format = payload.format.lower().strip()
        if doc_format not in ("pdf", "docx", "html", "txt"):
            doc_format = "pdf"

        # Unique file naming
        file_id = uuid.uuid4().hex[:12]
        safe_title = re.sub(r"[^a-zA-Z0-9_\-]", "_", payload.title.strip().lower())
        filename = f"{safe_title}_{file_id}.{doc_format}"
        file_path = os.path.join(DOCS_OUTPUT_DIR, filename)

        # Generate file artifact
        try:
            if doc_format == "pdf":
                self._generate_pdf(rendered_text, payload.title, file_path)
            elif doc_format == "docx":
                self._generate_docx(rendered_text, payload.title, file_path)
            elif doc_format == "html":
                self._generate_html(rendered_text, payload.title, file_path)
            else:
                self._generate_txt(rendered_text, file_path)
        except Exception as exc:
            logger.exception("Failed to build file artifact for format '%s': %s", doc_format, exc)
            # Fallback to plain text write
            with open(file_path, "w", encoding="utf-8") as f:
                f.write(rendered_text)

        # Normalize relative path with forward slashes for URLs
        rel_path = file_path.replace("\\", "/")

        doc_record = GeneratedDocument(
            company_id=company_id,
            template_id=template.id,
            title=payload.title,
            variables=payload.variables,
            format=doc_format,
            file_path_or_url=rel_path,
            generated_by=user_id,
        )
        self.session.add(doc_record)
        await self.session.commit()
        await self.session.refresh(doc_record)

        return GenerateDocumentResponse(
            documentId=str(doc_record.id),
            templateId=str(doc_record.template_id),
            title=doc_record.title,
            format=doc_record.format,
            filePathOrUrl=doc_record.file_path_or_url,
            variablesApplied=doc_record.variables or {},
            generatedAt=doc_record.created_at.isoformat(),
        )

    @staticmethod
    def _generate_pdf(content: str, title: str, output_path: str) -> None:
        """Render simple formatted PDF with reportlab."""
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

        c = canvas.Canvas(output_path, pagesize=letter)
        c.setFont("Helvetica-Bold", 16)
        c.drawString(50, 750, title)

        c.setFont("Helvetica", 11)
        y = 710
        for line in content.splitlines():
            if y < 60:
                c.showPage()
                c.setFont("Helvetica", 11)
                y = 740
            c.drawString(50, y, line[:90])
            y -= 16
        c.save()

    @staticmethod
    def _generate_docx(content: str, title: str, output_path: str) -> None:
        """Render DOCX file using python-docx."""
        import docx

        doc = docx.Document()
        doc.add_heading(title, level=1)
        for para in content.split("\n\n"):
            if para.strip():
                doc.add_paragraph(para.strip())
        doc.save(output_path)

    @staticmethod
    def _generate_html(content: str, title: str, output_path: str) -> None:
        """Render basic styled HTML file."""
        html_content = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>{title}</title>
    <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; line-height: 1.6; padding: 40px; color: #1e293b; max-width: 800px; margin: auto; }}
        h1 {{ color: #0f172a; border-bottom: 2px solid #e2e8f0; padding-bottom: 12px; }}
        p {{ margin-bottom: 16px; white-space: pre-wrap; }}
    </style>
</head>
<body>
    <h1>{title}</h1>
    <p>{content}</p>
</body>
</html>"""
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(html_content)

    @staticmethod
    def _generate_txt(content: str, output_path: str) -> None:
        """Write plain text file."""
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(content)
