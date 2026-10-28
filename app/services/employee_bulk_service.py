"""Tenant-scoped employee CSV/XLSX dry-run import and bulk export."""
from __future__ import annotations

import csv
import io
import re
import uuid
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.company import Branch
from app.models.department import Department
from app.models.employee import Employee
from app.services.tenant_onboarding_service import (
    TenantMappingError,
    validate_branch_department_manager,
    validate_employee_assignment,
)

EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
EXPORT_COLUMNS = [
    "employee_id", "first_name", "last_name", "personal_email", "company_email",
    "phone", "department", "designation", "joining_date", "employment_type",
    "employment_status", "status", "branch_id", "legal_entity_id", "department_id",
    "reporting_manager_id",
]


def _header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def parse_tabular_file(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    """Parse a CSV/XLSX upload into normalized dictionaries without DB writes."""
    lower = (filename or "").lower()
    if lower.endswith(".csv"):
        text = file_bytes.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise ValueError("The CSV file has no header row.")
        rows = [{_header(k): (v.strip() if isinstance(v, str) else v) for k, v in row.items() if k} for row in reader]
    elif lower.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        try:
            headers = [_header(value) for value in next(values)]
        except StopIteration as exc:
            raise ValueError("The XLSX file is empty.") from exc
        if not any(headers):
            raise ValueError("The XLSX file has no header row.")
        rows = []
        for raw in values:
            if not any(value is not None and value != "" for value in raw):
                continue
            rows.append({headers[index]: value for index, value in enumerate(raw) if index < len(headers) and headers[index]})
    else:
        raise ValueError("Only .csv and .xlsx files are supported.")
    return [row for row in rows if any(value is not None and value != "" for value in row.values())]


def _text(row: dict[str, Any], *keys: str, default: str = "") -> str:
    for key in keys:
        value = row.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return default


def _date(value: Any, field: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip()[:10])
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{field} must be an ISO date (YYYY-MM-DD).") from exc


def _uuid(value: Any, field: str) -> uuid.UUID | None:
    if value is None or not str(value).strip():
        return None
    try:
        return uuid.UUID(str(value).strip())
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{field} must be a valid UUID.") from exc


class EmployeeBulkService:
    def __init__(self, session: AsyncSession):
        self.session = session

    async def validate_employees(
        self, company_id: uuid.UUID, rows: list[dict[str, Any]]
    ) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        """Validate rows and manager links without mutating the database."""
        errors: list[dict[str, Any]] = []
        cleaned: list[dict[str, Any]] = []
        emails = [_text(row, "personal_email", "email").lower() for row in rows]
        codes = [_text(row, "employee_id", "employee_code", "code") for row in rows]
        seen_emails: set[str] = set()
        seen_codes: set[str] = set()

        existing = (await self.session.execute(select(Employee).where(Employee.company_id == company_id))).scalars().all()
        existing_by_code = {e.employee_id: e for e in existing}
        existing_manager_by_code = dict(existing_by_code)
        existing_emails = {e.personal_email.lower() for e in existing if e.personal_email}
        # personal_email is globally unique in the legacy schema; detect a
        # cross-tenant collision without revealing the other tenant's data.
        supplied_emails = {email for email in emails if email}
        if supplied_emails:
            global_duplicates = (await self.session.execute(
                select(Employee.personal_email).where(Employee.personal_email.in_(supplied_emails))
            )).scalars().all()
            existing_emails.update(email.lower() for email in global_duplicates)

        for index, row in enumerate(rows, start=2):
            row_errors: list[str] = []
            first_name = _text(row, "first_name", "firstname")
            last_name = _text(row, "last_name", "lastname", default="")
            full_name = _text(row, "name")
            if not first_name and full_name:
                first_name, *last_parts = full_name.split()
                last_name = last_name or " ".join(last_parts)
            email = _text(row, "personal_email", "email").lower()
            phone = _text(row, "phone", "mobile", "contact")
            code = _text(row, "employee_id", "employee_code", "code") or f"EMP-{uuid.uuid4().hex[:10].upper()}"
            if not first_name:
                row_errors.append("first_name is required.")
            if not email or not EMAIL_RE.match(email):
                row_errors.append("A valid personal_email is required.")
            if not phone:
                row_errors.append("phone is required.")
            if email in seen_emails or code in seen_codes:
                row_errors.append("Duplicate employee code or email within the file.")
            if email in existing_emails or code in existing_by_code:
                row_errors.append("Employee code or email already exists.")
            seen_emails.add(email)
            seen_codes.add(code)
            try:
                joining_date = _date(row.get("joining_date") or row.get("doj") or date.today(), "joining_date")
            except ValueError as exc:
                row_errors.append(str(exc))
                joining_date = date.today()
            try:
                branch_id = _uuid(row.get("branch_id"), "branch_id")
                legal_entity_id = _uuid(row.get("legal_entity_id"), "legal_entity_id")
                department_id = _uuid(row.get("department_id"), "department_id")
            except ValueError as exc:
                row_errors.append(str(exc))
                continue

            branch_code = _text(row, "branch_code")
            if branch_code and branch_id is None:
                branch = (await self.session.execute(select(Branch).where(
                    Branch.company_id == company_id, Branch.code == branch_code
                ))).scalar_one_or_none()
                if branch is None:
                    row_errors.append(f"Unknown branch_code: {branch_code}.")
                else:
                    branch_id = branch.id
            department_code = _text(row, "department_code")
            department_name = _text(row, "department", "department_name", default="General")
            if department_id is None and (department_code or department_name):
                stmt = select(Department).where(Department.company_id == company_id, Department.is_deleted == False)  # noqa: E712
                if department_code:
                    stmt = stmt.where(Department.department_code == department_code)
                else:
                    stmt = stmt.where(Department.department_name == department_name)
                department = (await self.session.execute(stmt)).scalar_one_or_none()
                if department is not None:
                    department_id = department.id
                    department_name = department.department_name
                elif department_code:
                    row_errors.append(f"Unknown department_code: {department_code}.")

            item = {
                "employee_id": code, "first_name": first_name, "last_name": last_name or "Employee",
                "personal_email": email, "company_email": _text(row, "company_email", default=email),
                "phone": phone, "department": department_name,
                "designation": _text(row, "designation", "title", default="Staff"),
                "joining_date": joining_date, "employment_type": _text(row, "employment_type", default="FULL_TIME").upper(),
                "employment_status": _text(row, "employment_status", default="PROBATION").upper(),
                "status": _text(row, "status", default="ACTIVE").upper(), "role": "employee",
                "branch_id": branch_id, "legal_entity_id": legal_entity_id, "department_id": department_id,
                "branch": branch_code or None, "work_location": _text(row, "work_location", "location") or None,
                "manager_code": _text(row, "reporting_manager_code", "manager_code", "reporting_manager_employee_id"),
            }
            try:
                item["legal_entity_id"] = await validate_branch_department_manager(
                    self.session, company_id=company_id, branch_id=branch_id,
                    legal_entity_id=legal_entity_id, department_id=department_id,
                )
            except TenantMappingError as exc:
                row_errors.append(str(exc))
            item["_row_number"] = index
            item["_manager_code"] = item["manager_code"]
            cleaned.append(item)
            errors.extend({"row": index, "employee_id": code, "message": message} for message in row_errors)

        error_rows = {error["row"] for error in errors}
        for item in cleaned:
            if item["_row_number"] in error_rows or not item["_manager_code"]:
                continue
            manager_code = item["_manager_code"]
            manager = existing_manager_by_code.get(manager_code)
            if manager is None:
                manager_item = next((candidate for candidate in cleaned if candidate["employee_id"] == manager_code), None)
                if manager_item is None or manager_item["_row_number"] in error_rows:
                    errors.append({"row": item["_row_number"], "message": f"Unknown reporting manager: {manager_code}."})
                    continue
                if manager_item["_row_number"] >= item["_row_number"]:
                    # Forward references are supported in the two-pass commit;
                    # only self references are inherently invalid.
                    if manager_item is item:
                        errors.append({"row": item["_row_number"], "message": "An employee cannot manage themselves."})
                        continue
            elif manager.id == item.get("id"):
                errors.append({"row": item["_row_number"], "message": "An employee cannot manage themselves."})
        valid = [item for item in cleaned if item["_row_number"] not in {e["row"] for e in errors}]
        return valid, errors

    async def import_employees(
        self, company_id: uuid.UUID, file_bytes: bytes, filename: str,
        *, dry_run: bool = False, actor_id: uuid.UUID | None = None,
    ) -> dict[str, Any]:
        rows = parse_tabular_file(file_bytes, filename)
        valid, errors = await self.validate_employees(company_id, rows)
        result = {
            "dry_run": dry_run, "total_processed": len(rows), "valid_rows": len(valid),
            "imported_count": 0, "skipped_count": len(rows) - len(valid), "errors": errors,
            "employees": [{k: (v.isoformat() if isinstance(v, date) else str(v) if isinstance(v, uuid.UUID) else v)
                           for k, v in item.items() if not k.startswith("_")} for item in valid] if dry_run else [],
        }
        if dry_run or errors or not valid:
            return result
        created_by_code: dict[str, Employee] = {}
        try:
            for item in valid:
                values = {k: v for k, v in item.items() if not k.startswith("_") and k != "manager_code"}
                employee = Employee(**values, company_id=company_id, created_by=actor_id, reporting_manager_id=None)
                self.session.add(employee)
                created_by_code[employee.employee_id] = employee
            await self.session.flush()
            existing = (await self.session.execute(select(Employee).where(Employee.company_id == company_id))).scalars().all()
            by_code = {e.employee_id: e for e in existing}
            by_code.update(created_by_code)
            for item in valid:
                employee = created_by_code[item["employee_id"]]
                if item["_manager_code"]:
                    manager = by_code.get(item["_manager_code"])
                    if manager is None or manager.id == employee.id:
                        raise TenantMappingError("Invalid reporting manager mapping.")
                    employee.reporting_manager_id = manager.id
                await validate_employee_assignment(self.session, employee)
            await self.session.commit()
        except Exception:
            await self.session.rollback()
            raise
        result["imported_count"] = len(valid)
        return result

    async def export_employees(self, company_id: uuid.UUID, file_format: str = "xlsx") -> tuple[bytes, str, str]:
        employees = (await self.session.execute(select(Employee).where(
            Employee.company_id == company_id, Employee.is_deleted == False  # noqa: E712
        ).order_by(Employee.employee_id))).scalars().all()
        rows = []
        for employee in employees:
            rows.append({column: (
                getattr(employee, column).isoformat() if isinstance(getattr(employee, column), (date, datetime))
                else str(getattr(employee, column)) if isinstance(getattr(employee, column), uuid.UUID)
                else getattr(employee, column) or ""
            ) for column in EXPORT_COLUMNS})
        file_format = file_format.lower()
        if file_format == "csv":
            output = io.StringIO()
            writer = csv.DictWriter(output, fieldnames=EXPORT_COLUMNS)
            writer.writeheader()
            writer.writerows(rows)
            return output.getvalue().encode("utf-8-sig"), "employees.csv", "text/csv; charset=utf-8"
        if file_format != "xlsx":
            raise ValueError("format must be csv or xlsx.")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Employees"
        sheet.append(EXPORT_COLUMNS)
        for row in rows:
            sheet.append([row[column] for column in EXPORT_COLUMNS])
        sheet.freeze_panes = "A2"
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue(), "employees.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


__all__ = ["EmployeeBulkService", "parse_tabular_file", "EXPORT_COLUMNS"]
