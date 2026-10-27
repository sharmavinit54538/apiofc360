"""E3 legacy payroll comparison with exact paise variances and persistence."""
from __future__ import annotations

import csv
import hashlib
import io
import uuid
from collections import Counter
from decimal import Decimal
from typing import Any, Iterable, Optional

from fastapi import HTTPException, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.employee import Employee
from app.models.payroll import PayrollRun, Payslip
from app.models.payroll_models import ParallelPayrollComparisonReport, PayrollRunEmployee
from app.services.payroll.money import as_decimal, paise_to_rupees, rupees_to_paise


class ParallelPayrollComparisonService:
    HEADS = ("gross", "basic", "pf", "esi", "pt", "lwf", "tds", "net_pay")
    REQUIRED_COLUMNS = ("employee_code", *HEADS)
    PF_CEILING_PAISE = 15000 * 100

    @staticmethod
    def _header(value: Any) -> str:
        return str(value or "").strip().lower().replace(" ", "_").replace("-", "_")

    @classmethod
    def parse_legacy_file(cls, content: bytes, filename: str = "legacy.csv") -> list[dict[str, Any]]:
        """Parse CSV/XLSX rows and normalize their headers without float math."""
        suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else "csv"
        if suffix in {"xlsx", "xlsm"}:
            try:
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                sheet = workbook.active
                values = list(sheet.iter_rows(values_only=True))
            except Exception as exc:
                raise HTTPException(status_code=422, detail=f"Invalid XLSX legacy payroll file: {exc}") from exc
            if not values:
                raise HTTPException(status_code=422, detail="Legacy payroll file is empty.")
            headers = [cls._header(v) for v in values[0]]
            raw_rows = [dict(zip(headers, row)) for row in values[1:] if any(v is not None for v in row)]
        elif suffix == "csv":
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise HTTPException(status_code=422, detail="CSV must be UTF-8 encoded.") from exc
            reader = csv.DictReader(io.StringIO(text))
            if not reader.fieldnames:
                raise HTTPException(status_code=422, detail="Legacy CSV has no header row.")
            headers = [cls._header(v) for v in reader.fieldnames]
            raw_rows = []
            for row in reader:
                normalized = {cls._header(k): v for k, v in row.items()}
                if any(str(v or "").strip() for v in normalized.values()):
                    raw_rows.append(normalized)
        else:
            raise HTTPException(status_code=422, detail="Only legacy CSV and XLSX files are supported.")

        missing = [c for c in cls.REQUIRED_COLUMNS if c not in headers]
        if missing:
            raise HTTPException(status_code=422, detail=f"Missing legacy payroll columns: {', '.join(missing)}")
        rows = []
        for line_number, raw in enumerate(raw_rows, start=2):
            code = str(raw.get("employee_code") or "").strip()
            if not code:
                raise HTTPException(status_code=422, detail=f"Missing employee_code on legacy row {line_number}.")
            normalized: dict[str, Any] = {"employee_code": code, "_row_number": line_number}
            for head in cls.HEADS:
                # An explicit *_paise column is accepted as a convenience for
                # API-to-API migrations; the documented columns are rupees.
                paise_key = f"{head}_paise"
                if paise_key in raw and raw.get(paise_key) not in (None, ""):
                    normalized[head] = int(as_decimal(str(raw[paise_key]).replace(",", "")))
                else:
                    value = str(raw.get(head) or "0").strip().replace(",", "").replace("₹", "").replace("Rs.", "")
                    try:
                        normalized[head] = rupees_to_paise(value or "0")
                    except Exception as exc:
                        raise HTTPException(status_code=422, detail=f"Invalid {head} amount on legacy row {line_number}.") from exc
            for optional in ("lop_days", "paid_days", "tax_regime", "regime", "pf_ceiling"):
                if optional in raw:
                    normalized[optional] = raw.get(optional)
            rows.append(normalized)
        return rows

    @staticmethod
    def _variance(payroll: int, legacy: int) -> dict[str, Any]:
        variance = int(payroll) - int(legacy)
        if legacy:
            percentage = float((Decimal(variance) / abs(Decimal(legacy)) * 100).quantize(Decimal("0.000001")))
        else:
            percentage = 0.0 if variance == 0 else None
        return {
            "legacy_paise": int(legacy),
            "payroll_paise": int(payroll),
            "variance_paise": variance,
            "deviation_percentage": percentage,
            "percentage_deviation": percentage,
            "legacy": float(paise_to_rupees(legacy)),
            "payroll": float(paise_to_rupees(payroll)),
        }

    @classmethod
    def _diagnose(cls, legacy: dict[str, Any], system: dict[str, Any], heads: dict[str, dict[str, Any]]) -> list[str]:
        diagnoses: list[str] = []
        nonzero = [h for h in cls.HEADS if heads[h]["variance_paise"]]
        if nonzero and all(abs(heads[h]["variance_paise"]) <= 100 for h in nonzero):
            diagnoses.append("ROUNDING_DIFF")
        legacy_lop = legacy.get("lop_days")
        if legacy_lop not in (None, "") and as_decimal(legacy_lop) != as_decimal(system.get("lop_days", 0)):
            diagnoses.append("LOP_DAYS_MISMATCH")
        legacy_paid = legacy.get("paid_days")
        if legacy_paid not in (None, "") and as_decimal(legacy_paid) != as_decimal(system.get("paid_days", 30)):
            if "LOP_DAYS_MISMATCH" not in diagnoses:
                diagnoses.append("LOP_DAYS_MISMATCH")
        pf_variance = heads["pf"]["variance_paise"]
        if pf_variance and (
            system.get("basic", 0) > cls.PF_CEILING_PAISE
            or legacy.get("basic", 0) > cls.PF_CEILING_PAISE
            or legacy.get("pf_ceiling") not in (None, "")
        ):
            diagnoses.append("PF_CEILING_DIFF")
        tds_variance = heads["tds"]["variance_paise"]
        if tds_variance and "ROUNDING_DIFF" not in diagnoses:
            legacy_regime = str(legacy.get("tax_regime") or legacy.get("regime") or "").upper()
            system_regime = str(system.get("tax_regime") or "").upper()
            if legacy_regime and system_regime and legacy_regime != system_regime:
                diagnoses.append("TAX_REGIME_DIFF")
            elif not legacy_regime or not system_regime or abs(tds_variance) > 100:
                # TDS has no other payroll-head explanation after rounding and
                # LOP checks; surface the regime as the first audit hypothesis.
                diagnoses.append("TAX_REGIME_DIFF")
        return diagnoses

    @classmethod
    async def compare_rows(
        cls,
        session: AsyncSession,
        *,
        run_id: uuid.UUID,
        legacy_rows: Iterable[dict[str, Any]],
        company_id: Optional[uuid.UUID] = None,
        legacy_file_name: Optional[str] = None,
        legacy_content: Optional[bytes] = None,
        created_by: Optional[uuid.UUID] = None,
        commit: bool = True,
    ) -> ParallelPayrollComparisonReport:
        run = await session.get(PayrollRun, run_id)
        if run is None or (company_id is not None and run.company_id != company_id):
            raise HTTPException(status_code=404, detail="Payroll run not found.")

        run_employees = (await session.execute(
            select(PayrollRunEmployee).where(PayrollRunEmployee.run_id == run_id)
        )).scalars().all()
        payslips = (await session.execute(select(Payslip).where(Payslip.payroll_run_id == run_id))).scalars().all()
        payslip_by_employee = {p.employee_id: p for p in payslips}
        employee_ids = [r.employee_id for r in run_employees]
        employees = (await session.execute(select(Employee).where(Employee.id.in_(employee_ids)))).scalars().all() if employee_ids else []
        employee_by_id = {e.id: e for e in employees}

        system_by_code: dict[str, dict[str, Any]] = {}
        for run_emp in run_employees:
            emp = employee_by_id.get(run_emp.employee_id)
            if emp is None:
                continue
            code = str(emp.employee_id or emp.id)
            ps = payslip_by_employee.get(run_emp.employee_id)
            earnings = run_emp.earnings_breakup or {}
            deductions = run_emp.deductions_breakup or {}
            def epaise(head: str, fallback: Any = 0) -> int:
                exact = earnings.get(f"{head}_paise")
                if exact is not None:
                    return int(exact)
                if head in earnings:
                    return rupees_to_paise(earnings[head])
                return rupees_to_paise(fallback)
            def dpaise(head: str, fallback: Any = 0) -> int:
                exact = deductions.get(f"{head}_paise")
                if exact is not None:
                    return int(exact)
                if head in deductions:
                    return rupees_to_paise(deductions[head])
                return rupees_to_paise(fallback)
            system_by_code[code] = {
                "employee_id": str(emp.id), "employee_code": code,
                "employee_name": f"{emp.first_name or ''} {emp.last_name or ''}".strip(),
                "gross": int(run_emp.gross_earnings_paise or 0),
                "basic": epaise("basic", ps.basic if ps else 0),
                "pf": dpaise("pf", ps.employee_pf if ps else 0),
                "esi": dpaise("esi", ps.employee_esi if ps else 0),
                "pt": dpaise("pt", ps.professional_tax if ps else 0),
                "lwf": dpaise("lwf", 0),
                "tds": dpaise("tds", ps.tds if ps else 0),
                "net_pay": int(run_emp.net_pay_paise or 0),
                "lop_days": str(run_emp.lop_days or 0), "paid_days": str(run_emp.paid_days or 30),
                "tax_regime": deductions.get("tax_regime"),
            }

        legacy_by_code = {str(row["employee_code"]): dict(row) for row in legacy_rows}
        employee_variances: list[dict[str, Any]] = []
        diagnoses_rows: list[dict[str, Any]] = []
        totals = {head: {"legacy": 0, "payroll": 0} for head in cls.HEADS}
        for code in sorted(set(system_by_code) | set(legacy_by_code)):
            system = system_by_code.get(code, {"employee_code": code, "employee_id": None, "employee_name": None, **{h: 0 for h in cls.HEADS}, "lop_days": 0, "paid_days": 30})
            legacy = legacy_by_code.get(code, {"employee_code": code, **{h: 0 for h in cls.HEADS}})
            heads = {head: cls._variance(system.get(head, 0), legacy.get(head, 0)) for head in cls.HEADS}
            diagnoses = cls._diagnose(legacy, system, heads)
            state = "MATCHED" if not any(heads[h]["variance_paise"] for h in cls.HEADS) else "VARIANCE"
            if code not in system_by_code:
                state = "MISSING_IN_PAYROLL"
            elif code not in legacy_by_code:
                state = "MISSING_IN_LEGACY"
            row = {
                "employee_code": code, "employee_id": system.get("employee_id"),
                "employee_name": system.get("employee_name"), "status": state,
                "heads": heads, "variances": heads, "diagnoses": diagnoses,
            }
            employee_variances.append(row)
            if diagnoses:
                diagnoses_rows.append({"employee_code": code, "codes": diagnoses})
            for head in cls.HEADS:
                totals[head]["legacy"] += int(legacy.get(head, 0))
                totals[head]["payroll"] += int(system.get(head, 0))

        head_variances = {
            head: cls._variance(values["payroll"], values["legacy"])
            for head, values in totals.items()
        }
        diagnosis_counts = Counter(code for row in diagnoses_rows for code in row["codes"])
        summary = {
            "run_id": str(run_id), "total_employees": len(employee_variances),
            "matched_employees": sum(r["status"] == "MATCHED" for r in employee_variances),
            "employees_with_variance": sum(r["status"] != "MATCHED" for r in employee_variances),
            "total_variance_paise": sum(abs(v["variance_paise"]) for v in head_variances.values()),
            "diagnosis_counts": dict(diagnosis_counts),
            "rounding": "ROUND_HALF_UP_NEAREST_RUPEE",
        }
        report = ParallelPayrollComparisonReport(
            id=uuid.uuid4(), company_id=run.company_id, run_id=run_id,
            legacy_file_name=legacy_file_name,
            legacy_file_sha256=hashlib.sha256(legacy_content).hexdigest() if legacy_content is not None else None,
            status="COMPLETED", employee_variances=employee_variances,
            head_variances=head_variances, diagnoses=diagnoses_rows, summary=summary,
            created_by=created_by,
        )
        session.add(report)
        if commit:
            await session.commit()
            await session.refresh(report)
        else:
            await session.flush()
        return report

    @classmethod
    async def compare_file(
        cls, session: AsyncSession, *, run_id: uuid.UUID, file: UploadFile,
        company_id: Optional[uuid.UUID] = None, created_by: Optional[uuid.UUID] = None,
    ) -> ParallelPayrollComparisonReport:
        content = await file.read()
        if len(content) > 20 * 1024 * 1024:
            raise HTTPException(status_code=413, detail="Legacy payroll file must be 20 MB or smaller.")
        rows = cls.parse_legacy_file(content, file.filename or "legacy.csv")
        return await cls.compare_rows(
            session, run_id=run_id, legacy_rows=rows, company_id=company_id,
            legacy_file_name=file.filename, legacy_content=content, created_by=created_by,
        )

    @classmethod
    async def get_report(cls, session: AsyncSession, report_id: uuid.UUID, company_id: Optional[uuid.UUID] = None):
        report = await session.get(ParallelPayrollComparisonReport, report_id)
        if report is None or (company_id is not None and report.company_id != company_id):
            raise HTTPException(status_code=404, detail="Parallel comparison report not found.")
        return report


ParallelComparisonService = ParallelPayrollComparisonService
